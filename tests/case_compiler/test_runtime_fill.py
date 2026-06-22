"""V3 上机回填：list_runtime_slots / apply_fills 的锁死 + 留空 + provenance 同步。

红线：用**合成 fixture**（emit 现造的 xlsx），不快照生产 KB 值当断言。
两条铁律对应的测试点：
- 不猜留空：apply_fills 给空值＝抽不出 → 该槽位如实留空（cell 不变、计 left_blank）。
- 不反复改：已填槽位不含 <RUNTIME> → 重填定位不到（not_found）、值绝不被覆盖。
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from main.case_compiler.provenance_ir import (
    CaseProvenance, StepIR, StepSource, RUNTIME_PLACEHOLDER, check_runtime_consistency,
)
from main.case_compiler.runtime_fill import list_runtime_slots, apply_fills
from main.ist_core.tools.device.emit_xlsx_tool import qa_emit_xlsx

_ROOT = Path(__file__).resolve().parents[2]


def _emit_fixture(out_name: str):
    """造一个含 2 个 <RUNTIME> 槽位(整值+部分模式) + 1 个确定值 check_point 的 case.xlsx。

    返回 (xlsx_path, provenance)。autoid==out_name（单 case，sidecar 落 outputs/<autoid>/）。
    """
    steps = [
        {"E": "APV_0", "F": "cmds_config", "G": "sdns listener 172.16.34.70 53"},
        {"E": "APV_0", "F": "cmd_config", "G": "show sdns service status"},   # 观测(整值前序)
        {"E": "check_point", "F": "found", "G": RUNTIME_PLACEHOLDER},          # 槽#0 整值不可知
        {"E": "APV_0", "F": "cmd_config", "G": "show statistics sdns pool"},  # 观测(部分前序)
        {"E": "check_point", "F": "found", "G": r"Hits:\s*<RUNTIME>"},         # 槽#1 部分模式
        {"E": "check_point", "F": "found", "G": "172.16.34.70"},               # 确定值,非槽
    ]
    prov = CaseProvenance(autoid=out_name, steps=[
        StepIR("APV_0", "cmds_config", "sdns listener 172.16.34.70 53", "G", StepSource("footprint", "sdns.listener")),
        StepIR("APV_0", "cmd_config", "show sdns service status", "G", StepSource("footprint", "sdns.service")),
        StepIR("check_point", "found", RUNTIME_PLACEHOLDER, "V", StepSource("device_runtime", "")),
        StepIR("APV_0", "cmd_config", "show statistics sdns pool", "G", StepSource("footprint", "sdns.pool")),
        StepIR("check_point", "found", r"Hits:\s*<RUNTIME>", "V", StepSource("device_runtime", "")),
        StepIR("check_point", "found", "172.16.34.70", "V", StepSource("manual", "10.5_cli:1")),
    ])
    r = qa_emit_xlsx.invoke({"autoid": out_name, "steps_json": json.dumps(steps),
                             "init_commands": "sdns on", "out_name": out_name,
                             "provenance_json": prov.to_json()})
    assert "已产出" in r, r
    return _ROOT / "workspace" / "outputs" / out_name / "case.xlsx", prov


def _g_at(xlsx_path, row):
    ws = openpyxl.load_workbook(str(xlsx_path), data_only=True).active
    return str(ws.cell(row, 7).value or "")


def _slots_by_g(xlsx):
    """按 current_g 取槽位（避免硬编行号；whole=整值占位, partial=部分模式）。"""
    slots = list_runtime_slots(xlsx)
    whole = next(s for s in slots if s.current_g == RUNTIME_PLACEHOLDER)
    partial = next(s for s in slots if s.current_g == r"Hits:\s*<RUNTIME>")
    return slots, whole, partial


def test_list_slots_finds_only_runtime_placeholders():
    xlsx, _ = _emit_fixture("t_fill_list")
    slots, whole, partial = _slots_by_g(xlsx)
    assert len(slots) == 2, [s.to_dict() for s in slots]
    # slot_id 行号化、稳定、互异；同 autoid 前缀
    assert whole.slot_id != partial.slot_id
    assert all(s.slot_id.startswith("t_fill_list#") for s in slots)
    # 前序观测命令被正确带出(回填值就从它的设备输出里抽)
    assert whole.observe_cmd == "show sdns service status"
    assert partial.observe_cmd == "show statistics sdns pool"


def test_apply_fills_whole_and_partial():
    xlsx, _ = _emit_fixture("t_fill_apply")
    _, whole, partial = _slots_by_g(xlsx)
    res = apply_fills(xlsx, [
        {"slot_id": whole.slot_id, "runtime_value": "active"},
        {"slot_id": partial.slot_id, "runtime_value": "42"},
    ], project_root=_ROOT, run_meta="build=test")
    assert set(res.filled) == {whole.slot_id, partial.slot_id}
    assert _g_at(xlsx, whole.row) == "active"              # 整值替换
    assert _g_at(xlsx, partial.row) == r"Hits:\s*42"       # 部分模式只换槽位
    # 填完即锁：再列槽位为空
    assert list_runtime_slots(xlsx) == []


def test_lock_no_overwrite_after_filled():
    """不反复改：已填槽位重填 → not_found，值不变；且不会错填到别的槽位。"""
    xlsx, _ = _emit_fixture("t_fill_lock")
    _, whole, partial = _slots_by_g(xlsx)
    apply_fills(xlsx, [{"slot_id": whole.slot_id, "runtime_value": "first"}], project_root=_ROOT)
    assert _g_at(xlsx, whole.row) == "first"
    # 用同一 slot_id 再填不同值 → 锁死，not_found，不覆盖
    res2 = apply_fills(xlsx, [{"slot_id": whole.slot_id, "runtime_value": "SECOND"}], project_root=_ROOT)
    assert res2.not_found == [whole.slot_id]
    assert res2.filled == []
    assert _g_at(xlsx, whole.row) == "first"            # 已填值绝不被改
    # 关键：填 whole 后，partial 的 id 没漂移、也没被错填，仍是占位
    assert RUNTIME_PLACEHOLDER in _g_at(xlsx, partial.row)


def test_empty_value_left_blank_not_guessed():
    """不猜留空：给空值＝抽不出 → 留空，cell 仍是占位。"""
    xlsx, _ = _emit_fixture("t_fill_blank")
    _, whole, _partial = _slots_by_g(xlsx)
    res = apply_fills(xlsx, [
        {"slot_id": whole.slot_id, "runtime_value": ""},        # 抽不出
        {"slot_id": whole.slot_id, "runtime_value": None},      # 同上(None)
    ], project_root=_ROOT)
    assert whole.slot_id in res.left_blank
    assert res.filled == []
    assert RUNTIME_PLACEHOLDER in _g_at(xlsx, whole.row)   # 仍是占位,没被猜
    # 留空的槽位仍可被后续真实回填(没锁死)
    assert any(s.slot_id == whole.slot_id for s in list_runtime_slots(xlsx))


def test_provenance_synced_to_device_verified():
    xlsx, _ = _emit_fixture("t_fill_prov")
    _, whole, _partial = _slots_by_g(xlsx)
    apply_fills(xlsx, [{"slot_id": whole.slot_id, "runtime_value": "active",
                        "evidence": "service status: active"}],
                project_root=_ROOT, run_meta="build=test;task=t1")
    sidecar = _ROOT / "workspace" / "outputs" / "t_fill_prov" / "case.provenance.json"
    prov = CaseProvenance.from_json(sidecar.read_text(encoding="utf-8"))
    # 被填的 check_point 来源转 device_verified，G 更新，且不再含占位
    cp0 = [s for s in prov.steps if s.E == "check_point"][0]
    assert cp0.source.kind == "device_verified"
    assert cp0.G == "active"
    assert "evidence" in cp0.source.ref
    # 仍含 <RUNTIME> 的部分模式槽位还是 device_runtime（没填）
    cp1 = [s for s in prov.steps if s.E == "check_point"][1]
    assert cp1.source.kind == "device_runtime" and RUNTIME_PLACEHOLDER in cp1.G
    # 同步后的 provenance 仍过一致性门（device_verified 无占位、device_runtime 有占位）
    assert check_runtime_consistency(prov) == []


def test_consistency_flags_verified_with_placeholder():
    """device_verified 却仍含占位（谎称已回填）→ 抓。"""
    prov = CaseProvenance(autoid="x", steps=[
        StepIR("APV_0", "cmd_config", "show sdns listener", "G", StepSource("footprint", "sdns.listener")),
        StepIR("check_point", "found", r"Hits:\s*<RUNTIME>", "V", StepSource("device_verified", "build=x")),
    ])
    probs = check_runtime_consistency(prov)
    assert len(probs) == 1 and "device_verified" in probs[0]


def test_consistency_partial_pattern_is_placeholder():
    """部分模式(前缀+<RUNTIME>)算占位：device_runtime 自洽、manual 被抓。"""
    ok = CaseProvenance(autoid="x", steps=[
        StepIR("check_point", "found", r"Hits:\s*<RUNTIME>", "V", StepSource("device_runtime", "")),
    ])
    assert check_runtime_consistency(ok) == []
    bad = CaseProvenance(autoid="x", steps=[
        StepIR("check_point", "found", r"Hits:\s*<RUNTIME>", "V", StepSource("manual", "x")),
    ])
    assert len(check_runtime_consistency(bad)) == 1
